"""Centro unificado de reportes de SICODE-UCT.

Los reportes trabajan únicamente con metadatos administrativos ya existentes en
SICODE. No almacenan ni exportan copias de expedientes o documentos sensibles.
"""

from __future__ import annotations

import csv
from datetime import datetime, time
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from flask import abort, make_response, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import or_

from app.models.alerta import Alerta
from app.models.bitacora import Bitacora
from app.models.expediente import Expediente
from app.models.prestamo import PrestamoExpediente
from app.routes.dashboard import dashboard_bp
from app.services.bitacora_service import registrar_bitacora


MAX_EXPORTACION = 5000
MAX_VISTA_PREVIA = 100

COLUMNAS = {
    "expedientes": [
        ("no_sp", "No. SP"),
        ("codigo_interno", "Código interno"),
        ("estado_administrativo", "Estado administrativo"),
        ("estado_documental", "Estado documental vigente"),
        ("disponibilidad", "Disponibilidad física"),
        ("folios_rectificados", "Folios rectificados"),
        ("anexos_rectificados", "Anexos rectificados"),
        ("ubicacion", "Ubicación física"),
        ("alertas_pendientes", "Alertas pendientes"),
        ("prestamo_activo", "Préstamo activo"),
        ("actualizado_en", "Última actualización"),
    ],
    "prestamos": [
        ("numero_control", "No. control"),
        ("no_sp", "No. SP"),
        ("estado", "Estado"),
        ("solicitante", "Solicitante"),
        ("persona_entrega", "Persona que entrega"),
        ("persona_recibe", "Persona que recibe"),
        ("fecha_prestamo", "Fecha préstamo"),
        ("fecha_estimada", "Devolución estimada"),
        ("fecha_real", "Devolución real"),
        ("observaciones", "Observaciones"),
    ],
    "alertas": [
        ("id", "ID"),
        ("no_sp", "No. SP"),
        ("tipo_alerta", "Tipo de alerta"),
        ("titulo", "Título"),
        ("gravedad", "Gravedad"),
        ("estado", "Estado"),
        ("origen", "Origen"),
        ("creado_en", "Creada"),
        ("cerrado_en", "Cerrada"),
        ("descripcion", "Descripción"),
    ],
    "bitacora": [
        ("id", "ID"),
        ("fecha", "Fecha Guatemala"),
        ("usuario", "Usuario"),
        ("accion", "Acción"),
        ("modulo", "Módulo"),
        ("entidad", "Entidad"),
        ("entidad_id", "Entidad ID"),
        ("no_sp", "No. SP"),
        ("motivo", "Motivo"),
        ("descripcion", "Descripción"),
    ],
    "consolidado": [
        ("no_sp", "No. SP"),
        ("codigo_interno", "Código interno"),
        ("estado_administrativo", "Estado administrativo"),
        ("estado_documental", "Estado documental vigente"),
        ("disponibilidad", "Disponibilidad física"),
        ("expediente_fisico", "Expediente físico"),
        ("folios_rectificados", "Folios rectificados"),
        ("anexos_rectificados", "Anexos rectificados"),
        ("documentos_principales", "Documentos cuerpo principal"),
        ("anexos_indice", "Anexos en índice"),
        ("alertas_pendientes", "Alertas pendientes"),
        ("prestamo_activo", "Préstamo activo"),
        ("ubicacion", "Ubicación física"),
        ("actualizado_en", "Última actualización"),
    ],
}

NOMBRES_REPORTES = {
    "expedientes": "Expedientes",
    "prestamos": "Préstamos y devoluciones",
    "alertas": "Alertas e incidentes",
    "bitacora": "Bitácora / auditoría",
    "consolidado": "Consolidado administrativo",
}


def _fecha(valor):
    return valor.strftime("%d/%m/%Y") if valor else ""


def _fecha_hora(valor):
    return valor.strftime("%d/%m/%Y %H:%M:%S") if valor else ""


def _parse_fecha(valor):
    try:
        return datetime.strptime(str(valor or ""), "%Y-%m-%d").date()
    except ValueError:
        return None


def _ubicacion_expediente(expediente):
    if not expediente.ubicaciones:
        return ""
    ubicacion = max(expediente.ubicaciones, key=lambda item: item.actualizado_en or item.creado_en or datetime.min)
    partes = [
        ("Archivador", ubicacion.archivador),
        ("SICOIN", ubicacion.sicoin),
        ("Estante", ubicacion.estante),
        ("Caja", ubicacion.caja),
        ("Módulo", ubicacion.modulo),
        ("Posición", ubicacion.posicion),
    ]
    return " · ".join(f"{etiqueta}: {valor}" for etiqueta, valor in partes if valor)


def _filtro_texto_expediente(consulta, busqueda):
    if not busqueda:
        return consulta
    patron = f"%{busqueda}%"
    return consulta.filter(or_(
        Expediente.no_sp.ilike(patron),
        Expediente.codigo_interno.ilike(patron),
        Expediente.nombre_referencia.ilike(patron),
    ))


def _filtrar_fecha(consulta, columna, desde, hasta):
    if desde:
        consulta = consulta.filter(columna >= datetime.combine(desde, time.min))
    if hasta:
        consulta = consulta.filter(columna <= datetime.combine(hasta, time.max))
    return consulta


def _filas_expedientes(busqueda, estado, desde, hasta, consolidado=False):
    consulta = Expediente.query.filter(Expediente.activo.is_(True))
    consulta = _filtro_texto_expediente(consulta, busqueda)
    consulta = _filtrar_fecha(consulta, Expediente.actualizado_en, desde, hasta)
    expedientes = consulta.order_by(Expediente.no_sp.asc()).limit(MAX_EXPORTACION).all()

    filas = []
    for expediente in expedientes:
        estado_documental = expediente.estado_fisico_documental
        if estado and estado not in {estado_documental, expediente.estado_administrativo, expediente.disponibilidad}:
            continue

        base = {
            "no_sp": expediente.no_sp,
            "codigo_interno": expediente.codigo_interno,
            "estado_administrativo": expediente.estado_administrativo,
            # Importante: usar la propiedad derivada vigente y no la expresión
            # SQL histórica de estado_fisico_documental.
            "estado_documental": estado_documental,
            "disponibilidad": expediente.disponibilidad,
            "folios_rectificados": expediente.folios_rectificados if expediente.folios_rectificados is not None else "",
            "anexos_rectificados": expediente.anexos_rectificados if expediente.anexos_rectificados is not None else "",
            "ubicacion": _ubicacion_expediente(expediente),
            "alertas_pendientes": len(expediente.alertas_pendientes),
            "prestamo_activo": expediente.prestamo_activo.numero_control if expediente.prestamo_activo else "",
            "actualizado_en": _fecha_hora(expediente.actualizado_en),
        }
        if consolidado:
            activos = expediente.documentos_activos
            base.update({
                "expediente_fisico": "Sí" if expediente.expediente_fisico_registrado else "No",
                "documentos_principales": sum(1 for documento in activos if not documento.es_anexo),
                "anexos_indice": sum(1 for documento in activos if documento.es_anexo),
            })
        filas.append(base)
    return filas


def _filas_prestamos(busqueda, estado, desde, hasta):
    consulta = PrestamoExpediente.query.join(Expediente)
    if busqueda:
        patron = f"%{busqueda}%"
        consulta = consulta.filter(or_(
            PrestamoExpediente.numero_control.ilike(patron),
            PrestamoExpediente.solicitante.ilike(patron),
            Expediente.no_sp.ilike(patron),
        ))
    if estado:
        consulta = consulta.filter(PrestamoExpediente.estado == estado)
    consulta = _filtrar_fecha(consulta, PrestamoExpediente.fecha_prestamo, desde, hasta)
    prestamos = consulta.order_by(PrestamoExpediente.fecha_prestamo.desc()).limit(MAX_EXPORTACION).all()
    return [
        {
            "numero_control": item.numero_control,
            "no_sp": item.expediente.no_sp if item.expediente else "",
            "estado": item.estado,
            "solicitante": item.solicitante,
            "persona_entrega": item.persona_entrega,
            "persona_recibe": item.persona_recibe,
            "fecha_prestamo": _fecha_hora(item.fecha_prestamo),
            "fecha_estimada": _fecha(item.fecha_estimada_devolucion),
            "fecha_real": _fecha_hora(item.fecha_real_devolucion),
            "observaciones": item.observaciones or "",
        }
        for item in prestamos
    ]


def _filas_alertas(busqueda, estado, desde, hasta):
    consulta = Alerta.query.join(Expediente)
    if busqueda:
        patron = f"%{busqueda}%"
        consulta = consulta.filter(or_(
            Expediente.no_sp.ilike(patron),
            Alerta.tipo_alerta.ilike(patron),
            Alerta.titulo.ilike(patron),
            Alerta.descripcion.ilike(patron),
        ))
    if estado:
        consulta = consulta.filter(Alerta.estado == estado)
    consulta = _filtrar_fecha(consulta, Alerta.creado_en, desde, hasta)
    alertas = consulta.order_by(Alerta.creado_en.desc()).limit(MAX_EXPORTACION).all()
    return [
        {
            "id": item.id,
            "no_sp": item.expediente.no_sp if item.expediente else "",
            "tipo_alerta": item.tipo_alerta,
            "titulo": item.titulo,
            "gravedad": item.gravedad,
            "estado": item.estado,
            "origen": item.origen,
            "creado_en": _fecha_hora(item.creado_en),
            "cerrado_en": _fecha_hora(item.cerrado_en),
            "descripcion": item.descripcion or "",
        }
        for item in alertas
    ]


def _filas_bitacora(busqueda, estado, desde, hasta):
    consulta = Bitacora.query
    accion = request.args.get("accion", "").strip()
    modulo = request.args.get("modulo", "").strip()
    usuario = request.args.get("usuario", "").strip()

    if busqueda:
        patron = f"%{busqueda}%"
        consulta = consulta.filter(or_(
            Bitacora.accion.ilike(patron),
            Bitacora.modulo.ilike(patron),
            Bitacora.descripcion.ilike(patron),
            Bitacora.entidad.ilike(patron),
            Bitacora.entidad_id.ilike(patron),
        ))
    if accion:
        consulta = consulta.filter(Bitacora.accion == accion)
    if modulo:
        consulta = consulta.filter(Bitacora.modulo == modulo)
    if usuario:
        consulta = consulta.join(Bitacora.usuario).filter_by(usuario=usuario)
    if estado:
        consulta = consulta.filter(Bitacora.modulo == estado)
    consulta = _filtrar_fecha(consulta, Bitacora.creado_en, desde, hasta)
    eventos = consulta.order_by(Bitacora.creado_en.desc()).limit(MAX_EXPORTACION).all()
    return [
        {
            "id": evento.id,
            "fecha": _fecha_hora(evento.creado_en_guatemala),
            "usuario": evento.usuario.usuario if evento.usuario else "Sistema / Sin usuario",
            "accion": evento.accion,
            "modulo": evento.modulo,
            "entidad": evento.entidad or "",
            "entidad_id": evento.entidad_id or "",
            "no_sp": evento.expediente.no_sp if evento.expediente else "",
            "motivo": evento.motivo or "",
            "descripcion": evento.descripcion or "",
        }
        for evento in eventos
    ]


def _validar_dataset(dataset):
    if dataset not in COLUMNAS:
        dataset = "expedientes"
    if dataset == "consolidado" and current_user.rol != "administrador":
        abort(403)
    return dataset


def _columnas_seleccionadas(dataset):
    permitidas = [clave for clave, _ in COLUMNAS[dataset]]
    recibidas = request.args.getlist("columnas")
    seleccion = [clave for clave in recibidas if clave in permitidas]
    if not seleccion:
        seleccion = permitidas[:8]
    return seleccion


def _obtener_filas(dataset):
    busqueda = request.args.get("q", "").strip()
    estado = request.args.get("estado", "").strip()
    desde = _parse_fecha(request.args.get("desde"))
    hasta = _parse_fecha(request.args.get("hasta"))

    if dataset == "expedientes":
        return _filas_expedientes(busqueda, estado, desde, hasta)
    if dataset == "consolidado":
        return _filas_expedientes(busqueda, estado, desde, hasta, consolidado=True)
    if dataset == "prestamos":
        return _filas_prestamos(busqueda, estado, desde, hasta)
    if dataset == "alertas":
        return _filas_alertas(busqueda, estado, desde, hasta)
    return _filas_bitacora(busqueda, estado, desde, hasta)


def _etiquetas(dataset):
    return dict(COLUMNAS[dataset])


def _valor_exportable(valor):
    """Evita que Excel/Calc interpreten texto controlado por usuarios como fórmula."""
    if not isinstance(valor, str):
        return valor
    visible = valor.lstrip(" \t\r\n")
    if visible.startswith(("=", "+", "-", "@")):
        return "'" + valor
    return valor


def _respuesta_sin_cache(respuesta):
    respuesta.headers["Cache-Control"] = "no-store, private"
    respuesta.headers["Pragma"] = "no-cache"
    respuesta.headers["X-Content-Type-Options"] = "nosniff"
    return respuesta


def _excel(dataset, columnas, filas):
    etiquetas = _etiquetas(dataset)
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte SICODE"
    ws.append([etiquetas[clave] for clave in columnas])
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for fila in filas:
        ws.append([_valor_exportable(fila.get(clave, "")) for clave in columnas])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for indice, clave in enumerate(columnas, start=1):
        ancho = min(max(len(etiquetas[clave]) + 2, 14), 45)
        ws.column_dimensions[ws.cell(row=1, column=indice).column_letter].width = ancho
        for celda in ws.iter_cols(min_col=indice, max_col=indice, min_row=2):
            for item in celda:
                item.alignment = Alignment(vertical="top", wrap_text=True)

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(
        salida,
        as_attachment=True,
        download_name=f"sicode_{dataset}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _csv(dataset, columnas, filas):
    etiquetas = _etiquetas(dataset)
    salida_texto = StringIO(newline="")
    escritor = csv.writer(salida_texto)
    escritor.writerow([etiquetas[clave] for clave in columnas])
    for fila in filas:
        escritor.writerow([_valor_exportable(fila.get(clave, "")) for clave in columnas])
    contenido = ("\ufeff" + salida_texto.getvalue()).encode("utf-8")
    return send_file(
        BytesIO(contenido),
        as_attachment=True,
        download_name=f"sicode_{dataset}.csv",
        mimetype="text/csv; charset=utf-8",
    )


def _pdf(dataset, columnas, filas):
    if len(columnas) > 10:
        abort(400, description="Para PDF seleccione un máximo de 10 columnas.")

    etiquetas = _etiquetas(dataset)
    salida = BytesIO()
    documento = SimpleDocTemplate(
        salida,
        pagesize=landscape(A4),
        rightMargin=22,
        leftMargin=22,
        topMargin=24,
        bottomMargin=24,
        title=f"SICODE-UCT · {NOMBRES_REPORTES[dataset]}",
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(f"<b>SICODE-UCT · {escape(NOMBRES_REPORTES[dataset])}</b>", estilos["Heading2"]),
        Paragraph(f"Registros exportados: {len(filas)}", estilos["BodyText"]),
        Spacer(1, 10),
    ]
    datos = [[Paragraph(f"<b>{escape(etiquetas[clave])}</b>", estilos["BodyText"]) for clave in columnas]]
    for fila in filas:
        datos.append([
            Paragraph(escape(str(fila.get(clave, "") or ""))[:900], estilos["BodyText"])
            for clave in columnas
        ])
    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    elementos.append(tabla)
    documento.build(elementos)
    salida.seek(0)
    return send_file(
        salida,
        as_attachment=True,
        download_name=f"sicode_{dataset}.pdf",
        mimetype="application/pdf",
    )


@dashboard_bp.before_app_request
def _redirigir_exportaciones_legacy():
    """Lleva exportaciones históricas al generador seguro sin romper sus URLs."""
    if not current_user.is_authenticated:
        return None

    if request.endpoint == "bitacora.exportar_excel":
        argumentos = request.args.to_dict(flat=True)
        argumentos.update(dataset="bitacora", formato="xlsx")
        return redirect(url_for("dashboard.reportes_exportar", **argumentos))

    if request.endpoint == "coordinacion_export.exportar_todos":
        if current_user.rol != "administrador":
            abort(403)
        return redirect(url_for("dashboard.reportes_exportar", dataset="consolidado", formato="xlsx"))
    return None


@dashboard_bp.route("/reportes")
@login_required
def reportes_centro():
    dataset = _validar_dataset(request.args.get("dataset", "expedientes"))
    columnas = _columnas_seleccionadas(dataset)
    filas = _obtener_filas(dataset)
    return render_template(
        "reportes/centro.html",
        dataset=dataset,
        datasets=NOMBRES_REPORTES,
        columnas_disponibles=COLUMNAS[dataset],
        columnas=columnas,
        etiquetas=_etiquetas(dataset),
        filas=filas[:MAX_VISTA_PREVIA],
        total=len(filas),
        limite_vista=MAX_VISTA_PREVIA,
        puede_exportar=not current_user.es_visor,
        es_admin=current_user.rol == "administrador",
    )


@dashboard_bp.route("/reportes/exportar")
@login_required
def reportes_exportar():
    dataset = _validar_dataset(request.args.get("dataset", "expedientes"))
    formato = request.args.get("formato", "xlsx").strip().lower()
    if formato not in {"xlsx", "csv", "pdf"}:
        abort(400)

    # El guard transversal de visor también bloquea endpoints cuyo nombre
    # contiene "exportar"; esta comprobación deja la política explícita aquí.
    if current_user.es_visor:
        abort(403)

    columnas = _columnas_seleccionadas(dataset)
    filas = _obtener_filas(dataset)
    registrar_bitacora(
        accion="EXPORTAR_CENTRO_REPORTES",
        modulo="Reportes",
        descripcion=f"Exportación {formato.upper()} del reporte {NOMBRES_REPORTES[dataset]}. Registros: {len(filas)}.",
        usuario_id=current_user.id,
        entidad="Reporte",
        entidad_id=dataset,
        datos_posteriores={
            "dataset": dataset,
            "formato": formato,
            "registros": len(filas),
            "columnas": columnas,
        },
    )

    if formato == "csv":
        respuesta = _csv(dataset, columnas, filas)
    elif formato == "pdf":
        respuesta = _pdf(dataset, columnas, filas)
    else:
        respuesta = _excel(dataset, columnas, filas)
    return _respuesta_sin_cache(respuesta)
