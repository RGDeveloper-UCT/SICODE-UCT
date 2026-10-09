"""Carga masiva con previsualización firmada y transacción única."""
import csv
import io
import re
from collections import defaultdict

from flask import abort, current_app, flash, redirect, render_template, request, url_for
from flask_login import current_user, login_required
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from openpyxl import load_workbook

from app import db
from app.models.documento_expediente import DocumentoExpediente
from app.models.expediente import Expediente
from app.routes.indice_documental import indice_documental_bp, _exigir_modificacion
from app.services.bitacora_service import registrar_bitacora

COLUMNAS = {"sp", "codigo_expediente", "seccion", "numero_anexo", "titulo_anexo",
            "orden", "nombre_documento", "folio_inicial", "folio_final", "estado_revision", "observaciones"}
ESTADOS = {"Pendiente de revisión", "Con observaciones", "Mal foliado", "Anexo pendiente"}
MAX_FILAS = 1500


class ImportacionInvalida(ValueError):
    pass


def _texto(v):
    return str(v).strip() if v is not None else ""


def _entero(v, campo, fila):
    valor = _texto(v)
    if not re.fullmatch(r"[0-9]+(?:\.0)?", valor):
        raise ImportacionInvalida(f"Fila {fila}: {campo} no es entero válido.")
    return int(float(valor))


def _leer(archivo, expediente):
    nombre = (archivo.filename or "").lower()
    contenido = archivo.read(1024 * 1024 + 1)
    if len(contenido) > 1024 * 1024:
        raise ImportacionInvalida("Máximo permitido: 1 MB.")
    if nombre.endswith(".csv"):
        try:
            filas = list(csv.DictReader(io.StringIO(contenido.decode("utf-8-sig"))))
        except (UnicodeError, csv.Error) as error:
            raise ImportacionInvalida("CSV inválido; utilice UTF-8.") from error
    elif nombre.endswith(".xlsx"):
        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
            hoja = libro["Indice"] if "Indice" in libro.sheetnames else libro.active
            valores = hoja.iter_rows(values_only=True)
            columnas = [_texto(x).lower() for x in next(valores)]
            filas = [dict(zip(columnas, valores_fila)) for valores_fila in valores if any(v is not None for v in valores_fila)]
            libro.close()
        except Exception as error:
            raise ImportacionInvalida("Excel inválido o sin hoja Indice.") from error
    else:
        raise ImportacionInvalida("Formato no permitido: utilice .xlsx o .csv.")
    if not filas or len(filas) > MAX_FILAS:
        raise ImportacionInvalida("Se requieren entre 1 y 1500 registros.")
    if not COLUMNAS.issubset(set(filas[0])):
        raise ImportacionInvalida("Columnas faltantes: " + ", ".join(sorted(COLUMNAS - set(filas[0]))))
    registros = []
    grupos = defaultdict(list)
    for fila, item in enumerate(filas, 2):
        sp = _entero(item.get("sp"), "sp", fila)
        if str(sp) != str(expediente.no_sp):
            raise ImportacionInvalida(f"Fila {fila}: SP {sp} distinto del expediente abierto.")
        codigo = _texto(item.get("codigo_expediente"))
        if codigo and codigo != expediente.codigo_interno:
            raise ImportacionInvalida(f"Fila {fila}: código interno no coincide.")
        seccion = _texto(item.get("seccion")).upper()
        if seccion not in ("PRINCIPAL", "ANEXO"):
            raise ImportacionInvalida(f"Fila {fila}: sección debe ser PRINCIPAL o ANEXO.")
        numero = _entero(item.get("numero_anexo"), "numero_anexo", fila) if seccion == "ANEXO" else None
        if numero is not None and not 1 <= numero <= 200:
            raise ImportacionInvalida(f"Fila {fila}: anexo fuera del rango 1-200.")
        if seccion == "PRINCIPAL" and _texto(item.get("numero_anexo")):
            raise ImportacionInvalida(f"Fila {fila}: sección principal no debe tener número de anexo.")
        inicio = _entero(item.get("folio_inicial"), "folio_inicial", fila)
        fin = _entero(item.get("folio_final"), "folio_final", fila)
        orden = _entero(item.get("orden"), "orden", fila)
        if inicio < 1 or fin < inicio or orden < 1:
            raise ImportacionInvalida(f"Fila {fila}: rango u orden incorrecto.")
        titulo = _texto(item.get("titulo_anexo")) if numero else ""
        nombre = _texto(item.get("nombre_documento"))
        estado = _texto(item.get("estado_revision")) or "Pendiente de revisión"
        observaciones = _texto(item.get("observaciones"))
        if not nombre or len(nombre) > 180 or len(titulo) > 180 or len(observaciones) > 1000:
            raise ImportacionInvalida(f"Fila {fila}: título, nombre u observaciones inválidos.")
        if estado not in ESTADOS:
            raise ImportacionInvalida(f"Fila {fila}: estado no admitido ({estado}).")
        registro = dict(seccion=seccion, numero=numero, titulo=titulo, orden=orden,
                        nombre=nombre, inicio=inicio, fin=fin, estado=estado, observaciones=observaciones)
        registros.append(registro)
        grupos[numero or 0].append(registro)
    alertas = []
    for numero, registros_grupo in grupos.items():
        vistos = set()
        titulo_anexo = {r["titulo"] for r in registros_grupo} if numero else set()
        if len(titulo_anexo) > 1:
            raise ImportacionInvalida(f"Anexo {numero}: títulos inconsistentes.")
        for r in sorted(registros_grupo, key=lambda x: (x["inicio"], x["fin"])):
            for folio in range(r["inicio"], r["fin"] + 1):
                if folio in vistos:
                    raise ImportacionInvalida(f"{'Principal' if numero == 0 else 'Anexo '+str(numero)}: traslape en folio {folio}.")
                vistos.add(folio)
        if vistos and (min(vistos) != 1 or len(vistos) != max(vistos)):
            alertas.append(f"{'Principal' if numero == 0 else 'Anexo '+str(numero)}: faltan folios en la secuencia.")
    return registros, alertas


def _firmador():
    return URLSafeTimedSerializer(current_app.secret_key, salt="indice-documental-masivo-v1")


@indice_documental_bp.route("/expedientes/<int:expediente_id>/indice-documental/importar", methods=["GET", "POST"])
@login_required
def importar_indice(expediente_id):
    _exigir_modificacion()
    expediente = Expediente.query.get_or_404(expediente_id)
    if not expediente.expediente_fisico_registrado:
        abort(403)
    registros, alertas, token = [], [], None
    if request.method == "POST":
        try:
            if request.form.get("accion") == "confirmar":
                try:
                    paquete = _firmador().loads(request.form.get("token", ""), max_age=1800)
                except (BadSignature, SignatureExpired) as error:
                    raise ImportacionInvalida("La previsualización expiró o fue alterada. Suba el archivo nuevamente.") from error
                if paquete.get("expediente_id") != expediente.id or paquete.get("usuario_id") != current_user.id:
                    abort(403)
                registros = paquete["registros"]
                alertas = paquete["alertas"]
                if alertas:
                    raise ImportacionInvalida("Hay saltos de foliación. Corrija el archivo y vuelva a cargarlo.")
                # No modificar índices existentes. Evita duplicados, relaciones ambiguas e importaciones repetidas.
                if DocumentoExpediente.query.filter_by(expediente_id=expediente.id, activo=True).first():
                    raise ImportacionInvalida("Este SP ya tiene índice activo. La carga inicial no sobrescribe registros: revise el índice antes de continuar.")
                for r in registros:
                    documento = DocumentoExpediente(
                        expediente_id=expediente.id, nombre_documento=r["nombre"],
                        tipo_documento="Anexo" if r["numero"] else "Documento",
                        folio_inicio=r["inicio"], folio_fin=r["fin"],
                        total_folios=r["fin"] - r["inicio"] + 1,
                        estado_revision=r["estado"], es_anexo=bool(r["numero"]),
                        indice_numero_anexo=r["numero"], indice_titulo_anexo=r["titulo"] or None,
                        indice_orden=r["orden"], observaciones=r["observaciones"] or None, activo=True)
                    db.session.add(documento)
                registrar_bitacora(
                    accion="IMPORTAR_INDICE_DOCUMENTAL", modulo="Índice documental",
                    descripcion=f"Importación inicial de {len(registros)} documentos para SP {expediente.no_sp}.",
                    usuario_id=current_user.id, expediente_id=expediente.id,
                    entidad="Expediente", entidad_id=expediente.id,
                    datos_posteriores={"total": len(registros), "anexos": sorted(set(r["numero"] for r in registros if r["numero"]))},
                    commit=False)
                db.session.commit()
                flash(f"Índice importado: {len(registros)} registros. Pendientes de revisión física.", "success")
                return redirect(url_for("indice_documental.listado", expediente_id=expediente.id))
            archivo = request.files.get("archivo")
            if not archivo:
                raise ImportacionInvalida("Seleccione un Excel o CSV.")
            registros, alertas = _leer(archivo, expediente)
            if DocumentoExpediente.query.filter_by(expediente_id=expediente.id, activo=True).first():
                alertas.append("El SP ya tiene documentos activos. No se permite la importación inicial.")
            token = _firmador().dumps(dict(expediente_id=expediente.id, usuario_id=current_user.id,
                                          registros=registros, alertas=alertas))
        except ImportacionInvalida as error:
            flash(str(error), "danger")
        except Exception:
            db.session.rollback()
            raise
    return render_template("indice_documental/importar.html", expediente=expediente,
                           registros=registros, alertas=alertas, token=token)
