"""Conciliación conservadora de índices File Server contra los metadatos de SICODE.

No lee ni almacena documentos completos, no modifica datos manuales ni declara
rectificaciones físicas. El Excel es fuente de contraste, no fuente de verdad.
"""
from __future__ import annotations

import io
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy.orm import joinedload

from app.models.anexo_rectificado import AnexoRectificado
from app.models.coordinacion import AnexoCoordinacion, RegistroCoordinacion
from app.models.documento_expediente import DocumentoExpediente
from app.models.expediente import Expediente
from app.services.catalogo_anexos_service import catalogo_plano

COLUMNAS = (
    "sp", "codigo_expediente", "seccion", "numero_anexo", "titulo_anexo",
    "orden", "nombre_documento", "tipo_documento", "folio_inicial",
    "folio_final", "estado_revision", "fecha_recepcion", "rc", "providencia",
    "tipo_anexo", "es_vencido", "observaciones", "referencia_file_server",
)
OBLIGATORIAS = {"sp", "seccion", "orden", "nombre_documento"}
TIPOS = {"Documento", "Anexo", "Oficio", "Resolución", "Acta", "Informe", "Otro"}
ESTADOS = {"Pendiente de revisión", "Verificado", "Con observaciones",
           "Mal foliado", "Anexo pendiente"}
MAX_FILAS = 12000
MAX_BYTES = 8 * 1024 * 1024
ESTADOS_NO_IMPORTABLES = {
    "SP_NO_EXISTE", "CODIGO_DIFERENTE", "SIN_EXPEDIENTE_FISICO",
    "SIN_FOLIACION", "DUPLICADO_EXCEL", "DIFERENCIA_INDICE",
    "DIFERENCIA_ANEXO", "ANEXO_AMBIGUO", "REGISTRO_ANULADO",
}
ESTADOS_EXPLICADOS = {
    "NUEVO": "Se incorporará al índice, pendiente de revisión física",
    "COINCIDE": "Ya existe en SICODE; se conserva sin cambios",
    "SP_NO_EXISTE": "No existe SP maestro; no se crea automáticamente",
    "CODIGO_DIFERENTE": "El código interno no coincide con el SP",
    "SIN_EXPEDIENTE_FISICO": "No hay expediente físico registrado",
    "SIN_FOLIACION": "Faltan folios; se requieren datos físicos",
    "DUPLICADO_EXCEL": "Registro repetido o con folios cruzados en el Excel",
    "DIFERENCIA_INDICE": "Contradicción con el índice ya existente",
    "DIFERENCIA_ANEXO": "Discrepancia con anexos de Coordinación/rectificados",
    "ANEXO_AMBIGUO": "No se puede identificar el anexo sin riesgo",
    "REGISTRO_ANULADO": "Existe antecedente anulado; revisión antes de recrear",
    "SOLO_SICODE": "Registro manual existente que no figura en el Excel recibido; se conserva",
}


class IndiceFileServerInvalido(ValueError):
    pass


def texto(v, maximo=500):
    if v is None:
        return ""
    resultado = str(v).strip()
    if len(resultado) > maximo:
        raise IndiceFileServerInvalido(f"Valor excede {maximo} caracteres.")
    return resultado


def normalizar(v):
    s = unicodedata.normalize("NFKD", texto(v))
    s = "".join(x for x in s if not unicodedata.combining(x))
    return " ".join(re.sub(r"[^a-zA-Z0-9]+", " ", s.casefold()).split())


def normalizar_sp(v):
    s = texto(v, 50).upper()
    s = re.sub(r"^SP\s*[-:#]?\s*", "", s, flags=re.I)
    if re.fullmatch(r"\d+(?:\.0)?", s):
        return str(int(s[:-2] if s.endswith(".0") else s))
    return s


def entero(v, campo, fila, obligatorio=True):
    s = texto(v, 30)
    if not s and not obligatorio:
        return None
    if not re.fullmatch(r"\d+(?:\.0)?", s):
        raise IndiceFileServerInvalido(f"Fila {fila}: {campo} debe ser entero.")
    return int(s[:-2] if s.endswith(".0") else s)


def fecha_iso(v, fila):
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    try:
        return datetime.strptime(str(v).strip(), "%Y-%m-%d").date().isoformat()
    except ValueError as error:
        raise IndiceFileServerInvalido(
            f"Fila {fila}: fecha_recepcion debe tener formato AAAA-MM-DD."
        ) from error


def leer_excel(archivo):
    nombre = (archivo.filename or "").lower()
    if not nombre.endswith(".xlsx"):
        raise IndiceFileServerInvalido("Cargue únicamente el formato Excel .xlsx.")
    contenido = archivo.stream.read(MAX_BYTES + 1)
    if len(contenido) > MAX_BYTES:
        raise IndiceFileServerInvalido("El Excel no puede superar 8 MB.")
    try:
        with ZipFile(io.BytesIO(contenido)) as comprimido:
            piezas = comprimido.infolist()
            if len(piezas) > 1000 or sum(p.file_size for p in piezas) > 80 * 1024 * 1024:
                raise IndiceFileServerInvalido(
                    "Excel demasiado grande al descomprimirse; divida la importación.")
        libro = load_workbook(io.BytesIO(contenido), read_only=True,
                              data_only=True, keep_links=False)
        if "Indice" not in libro.sheetnames:
            raise IndiceFileServerInvalido("Falta la hoja Indice.")
        hoja = libro["Indice"]
        if hoja.max_row and hoja.max_row > MAX_FILAS + 1:
            raise IndiceFileServerInvalido("El máximo es de 12000 filas de datos.")
        iterador = hoja.iter_rows(values_only=True)
        cabecera = [texto(c, 100).lower() for c in next(iterador)]
        if len(cabecera) != len(set(cabecera)):
            raise IndiceFileServerInvalido("La hoja contiene encabezados repetidos.")
        faltantes = OBLIGATORIAS - set(cabecera)
        if faltantes:
            raise IndiceFileServerInvalido(
                "Columnas obligatorias faltantes: " + ", ".join(sorted(faltantes))
            )
        filas = []
        for numero_fila, valores in enumerate(iterador, 2):
            if not any(v is not None and str(v).strip() for v in valores):
                continue
            if len(filas) >= MAX_FILAS:
                raise IndiceFileServerInvalido("Máximo 12000 filas en un lote.")
            bruto = dict(zip(cabecera, valores))
            sp = normalizar_sp(bruto.get("sp"))
            if not sp:
                raise IndiceFileServerInvalido(f"Fila {numero_fila}: falta sp.")
            seccion = texto(bruto.get("seccion"), 20).upper()
            if seccion not in {"PRINCIPAL", "ANEXO"}:
                raise IndiceFileServerInvalido(
                    f"Fila {numero_fila}: seccion debe ser PRINCIPAL o ANEXO."
                )
            numero_anexo = entero(bruto.get("numero_anexo"), "numero_anexo",
                                  numero_fila, obligatorio=seccion == "ANEXO")
            if seccion == "PRINCIPAL" and numero_anexo is not None:
                raise IndiceFileServerInvalido(
                    f"Fila {numero_fila}: PRINCIPAL no lleva numero_anexo."
                )
            if numero_anexo is not None and not 1 <= numero_anexo <= 200:
                raise IndiceFileServerInvalido(
                    f"Fila {numero_fila}: numero_anexo debe estar entre 1 y 200."
                )
            orden = entero(bruto.get("orden"), "orden", numero_fila)
            if orden < 1:
                raise IndiceFileServerInvalido(f"Fila {numero_fila}: orden inválido.")
            nombre_documento = texto(bruto.get("nombre_documento"), 180)
            if not nombre_documento:
                raise IndiceFileServerInvalido(f"Fila {numero_fila}: falta nombre_documento.")
            inicio = entero(bruto.get("folio_inicial"), "folio_inicial", numero_fila, False)
            fin = entero(bruto.get("folio_final"), "folio_final", numero_fila, False)
            if (inicio is None) != (fin is None) or (
                inicio is not None and (inicio < 1 or fin < inicio)
            ):
                raise IndiceFileServerInvalido(
                    f"Fila {numero_fila}: folios incompletos o rango inválido."
                )
            tipo = texto(bruto.get("tipo_documento"), 80) or (
                "Anexo" if seccion == "ANEXO" else "Documento"
            )
            if tipo not in TIPOS:
                raise IndiceFileServerInvalido(f"Fila {numero_fila}: tipo_documento desconocido.")
            estado = texto(bruto.get("estado_revision"), 80) or "Pendiente de revisión"
            if estado not in ESTADOS:
                raise IndiceFileServerInvalido(f"Fila {numero_fila}: estado_revision inválido.")
            vencido = texto(bruto.get("es_vencido"), 15).upper() or "DESCONOCIDO"
            if vencido not in {"SI", "NO", "DESCONOCIDO"}:
                raise IndiceFileServerInvalido(
                    f"Fila {numero_fila}: es_vencido debe ser SI, NO o DESCONOCIDO."
                )
            filas.append({
                "fila": numero_fila, "sp": sp,
                "codigo_expediente": texto(bruto.get("codigo_expediente"), 50),
                "seccion": seccion, "numero_anexo": numero_anexo,
                "titulo_anexo": texto(bruto.get("titulo_anexo"), 180),
                "orden": orden, "nombre_documento": nombre_documento,
                "tipo_documento": tipo, "folio_inicial": inicio,
                "folio_final": fin, "estado_revision": estado,
                "fecha_recepcion": fecha_iso(bruto.get("fecha_recepcion"), numero_fila),
                "rc": texto(bruto.get("rc"), 80),
                "providencia": texto(bruto.get("providencia"), 120),
                "tipo_anexo": texto(bruto.get("tipo_anexo"), 120),
                "es_vencido": vencido,
                "observaciones": texto(bruto.get("observaciones"), 1000),
                "referencia_file_server": texto(bruto.get("referencia_file_server"), 160),
            })
        libro.close()
    except IndiceFileServerInvalido:
        raise
    except Exception as error:
        raise IndiceFileServerInvalido("Archivo Excel dañado o inválido.") from error
    if not filas:
        raise IndiceFileServerInvalido("No hay registros en la hoja Indice.")
    return filas


def normalizar_tipo_anexo(valor):
    dato = normalizar(valor)
    for codigo, ficha in catalogo_plano().items():
        if dato in {normalizar(codigo), normalizar(ficha["titulo"])}:
            return codigo
    return dato


def numero_existente(doc):
    if not doc.es_anexo:
        return 0
    if doc.indice_numero_anexo is not None:
        return int(doc.indice_numero_anexo)
    anexo = doc.anexo_recepcion
    if anexo and str(anexo.numero_anexo or "").strip().isdigit():
        return int(anexo.numero_anexo)
    return None


def _coincidencia_anexo(fila, anexos, rectificados):
    """Exige no contradicción con la recepción y la rectificación existente."""
    candidatos = [a for a in anexos if str(a.numero_anexo or "").strip().isdigit()
                  and int(a.numero_anexo) == fila["numero_anexo"]]
    if len(candidatos) > 1:
        return "ANEXO_AMBIGUO", None, "El número aparece en varias recepciones."
    anexo = candidatos[0] if candidatos else None
    if anexo:
        recepcion = anexo.registro
        comparaciones = (
            (fila["titulo_anexo"], anexo.titulo),
            (fila["rc"], recepcion.rc),
            (fila["providencia"], recepcion.providencia),
            (fila["fecha_recepcion"],
             recepcion.fecha_recepcion.isoformat() if recepcion.fecha_recepcion else ""),
        )
        for entrante, actual in comparaciones:
            if entrante and actual and normalizar(entrante) != normalizar(actual):
                return "DIFERENCIA_ANEXO", anexo, "Metadatos distintos de Coordinación."
        if (fila["tipo_anexo"] and anexo.tipo_anexo and
                normalizar_tipo_anexo(fila["tipo_anexo"]) !=
                normalizar_tipo_anexo(anexo.tipo_anexo)):
            return "DIFERENCIA_ANEXO", anexo, "Tipo de anexo distinto de Coordinación."
        if fila["es_vencido"] != "DESCONOCIDO" and bool(anexo.es_vencido) != (
            fila["es_vencido"] == "SI"
        ):
            return "DIFERENCIA_ANEXO", anexo, "Vigencia/vencimiento no coincide."
    rect = [a for a in rectificados if str(a.numero_anexo or "").strip().isdigit()
            and int(a.numero_anexo) == fila["numero_anexo"] and a.activo]
    for r in rect:
        if fila["titulo_anexo"] and r.titulo and normalizar(fila["titulo_anexo"]) != normalizar(r.titulo):
            return "DIFERENCIA_ANEXO", anexo, "Título distinto de la rectificación física."
    return None, anexo, ""


def planificar(filas):
    """Produce decisiones reproducibles, sin cambiar la sesión ni la base."""
    sp_ids = set(fila["sp"] for fila in filas)
    expedientes = {e.no_sp: e for e in Expediente.query.filter(
        Expediente.no_sp.in_(sp_ids)).all()}
    ids = [e.id for e in expedientes.values()]
    documentos = defaultdict(list)
    anexos = defaultdict(list)
    rectificados = defaultdict(list)
    if ids:
        for d in DocumentoExpediente.query.filter(
            DocumentoExpediente.expediente_id.in_(ids)
        ).options(joinedload(DocumentoExpediente.anexo_recepcion)).all():
            documentos[d.expediente_id].append(d)
        for a in AnexoCoordinacion.query.join(RegistroCoordinacion).filter(
            RegistroCoordinacion.expediente_id.in_(ids)
        ).options(joinedload(AnexoCoordinacion.registro),
                  joinedload(AnexoCoordinacion.documento_expediente)).all():
            anexos[a.registro.expediente_id].append(a)
        for a in AnexoRectificado.query.filter(AnexoRectificado.expediente_id.in_(ids)).all():
            rectificados[a.expediente_id].append(a)

    decisiones = []
    vistos = set()
    propuestos = defaultdict(list)
    cotejados = set()
    for r in filas:
        resultado = dict(r)
        resultado.update(estado="NUEVO", detalle="", expediente_id=None,
                         anexo_coordinacion_id=None)
        expediente = expedientes.get(r["sp"])
        if not expediente:
            resultado.update(estado="SP_NO_EXISTE", detalle="No se creará un nuevo SP.")
        elif r["codigo_expediente"] and r["codigo_expediente"] != expediente.codigo_interno:
            resultado.update(estado="CODIGO_DIFERENTE", detalle=expediente.codigo_interno)
        elif not expediente.expediente_fisico_registrado:
            resultado.update(estado="SIN_EXPEDIENTE_FISICO")
        else:
            resultado["expediente_id"] = expediente.id
            ambito = (expediente.id, r["seccion"], r["numero_anexo"] or 0)
            identidad = (
                ambito, normalizar(r["nombre_documento"]),
                r["folio_inicial"], r["folio_final"],
            )
            if identidad in vistos:
                resultado.update(estado="DUPLICADO_EXCEL", detalle="Fila repetida.")
            elif r["folio_inicial"] is None:
                resultado.update(estado="SIN_FOLIACION")
            else:
                conflicto, anexo, detalle = (None, None, "")
                if r["seccion"] == "ANEXO":
                    conflicto, anexo, detalle = _coincidencia_anexo(
                        r, anexos[expediente.id], rectificados[expediente.id]
                    )
                    if anexo:
                        resultado["anexo_coordinacion_id"] = anexo.id
                mismos_ambitos = [d for d in documentos[expediente.id]
                                  if (0 if not d.es_anexo else numero_existente(d)) ==
                                  (r["numero_anexo"] or 0)
                                  and bool(d.es_anexo) == (r["seccion"] == "ANEXO")]
                exactos = [d for d in mismos_ambitos if (
                    normalizar(d.nombre_documento) == normalizar(r["nombre_documento"])
                    and d.folio_inicio == r["folio_inicial"]
                    and d.folio_fin == r["folio_final"])]
                if conflicto:
                    resultado.update(estado=conflicto, detalle=detalle)
                elif exactos:
                    cotejados.update(d.id for d in exactos)
                    if any(d.activo for d in exactos):
                        resultado.update(estado="COINCIDE", detalle="Registro ya activo; sin cambios.")
                    else:
                        resultado.update(estado="REGISTRO_ANULADO")
                elif any(d.activo and (
                    normalizar(d.nombre_documento) == normalizar(r["nombre_documento"])
                    or d.folio_inicio <= r["folio_final"] and d.folio_fin >= r["folio_inicial"]
                ) for d in mismos_ambitos):
                    cotejados.update(d.id for d in mismos_ambitos if d.activo and (
                        normalizar(d.nombre_documento) == normalizar(r["nombre_documento"])
                        or d.folio_inicio <= r["folio_final"] and d.folio_fin >= r["folio_inicial"]
                    ))
                    resultado.update(estado="DIFERENCIA_INDICE",
                                     detalle="Nombre o rango de folios en conflicto.")
                elif any(d.activo and d.es_anexo and numero_existente(d) is None
                         and normalizar(d.nombre_documento) == normalizar(r["nombre_documento"])
                         for d in documentos[expediente.id]) and r["seccion"] == "ANEXO":
                    resultado.update(estado="ANEXO_AMBIGUO",
                                     detalle="Hay anexo manual sin número que puede ser el mismo.")
                elif any(p["folio_inicial"] <= r["folio_final"] and
                         p["folio_final"] >= r["folio_inicial"]
                         for p in propuestos[ambito]):
                    resultado.update(estado="DUPLICADO_EXCEL",
                                     detalle="Rango superpuesto en este mismo ámbito.")
                else:
                    propuestos[ambito].append(r)
                    if r["seccion"] == "ANEXO" and anexo is None:
                        resultado["detalle"] = (
                            "No hay recepción de Coordinación asociable; "
                            "se creará solo el metadato de índice, sin simular recepción."
                        )
                    if r["seccion"] == "PRINCIPAL" and expediente.folios_rectificados is not None:
                        if r["folio_final"] > expediente.folios_rectificados:
                            resultado["detalle"] = (
                                "Excede folios físicamente rectificados; revisar posteriormente."
                            )
            vistos.add(identidad)
        decisiones.append(resultado)

    # Un rango repetido o cruzado en el propio Excel no puede importarse
    # parcialmente. Marcar también la PRIMERA fila como conflicto evita que
    # el orden del archivo cambie qué documento se inserta.
    grupos_excel = defaultdict(list)
    for registro in decisiones:
        if isinstance(registro["fila"], int):
            grupos_excel[(
                registro["sp"], registro["seccion"],
                registro["numero_anexo"] or 0
            )].append(registro)
    for grupo in grupos_excel.values():
        for i, candidato in enumerate(grupo):
            if candidato["estado"] != "NUEVO":
                continue
            if any(
                i != j and otro["folio_inicial"] is not None
                and candidato["folio_inicial"] <= otro["folio_final"]
                and candidato["folio_final"] >= otro["folio_inicial"]
                for j, otro in enumerate(grupo)
            ):
                candidato.update(
                    estado="DUPLICADO_EXCEL",
                    detalle="El Excel contiene otro registro con folios superpuestos "
                            "en el mismo SP y ámbito; se bloquean ambos.",
                )

    # Conciliación en ambas direcciones: lo que está en SICODE y no se encuentra
    # representado en el Excel de los SP cargados NO se elimina ni se anula.
    for sp in sorted(sp_ids):
        expediente = expedientes.get(sp)
        if expediente is None:
            continue
        for documento in documentos[expediente.id]:
            if not documento.activo or documento.id in cotejados:
                continue
            ambito = "ANEXO" if documento.es_anexo else "PRINCIPAL"
            numero = numero_existente(documento)
            datos_mismos = [r for r in filas if r["sp"] == sp and
                           r["seccion"] == ambito and
                           (r["numero_anexo"] or 0) == (numero or 0) and
                           normalizar(r["nombre_documento"]) ==
                           normalizar(documento.nombre_documento)]
            if datos_mismos:
                continue
            decisiones.append({
                "fila": f"SICODE-{documento.id}", "sp": sp,
                "codigo_expediente": expediente.codigo_interno,
                "seccion": ambito, "numero_anexo": numero,
                "titulo_anexo": documento.indice_titulo_anexo or "",
                "orden": documento.indice_orden or 0,
                "nombre_documento": documento.nombre_documento,
                "tipo_documento": documento.tipo_documento,
                "folio_inicial": documento.folio_inicio,
                "folio_final": documento.folio_fin,
                "estado_revision": documento.estado_revision,
                "fecha_recepcion": "", "rc": "", "providencia": "",
                "tipo_anexo": "", "es_vencido": "DESCONOCIDO",
                "observaciones": "", "referencia_file_server": "",
                "estado": "SOLO_SICODE",
                "detalle": "Existe en SICODE; no aparece en el Excel de este SP.",
                "expediente_id": expediente.id,
                "anexo_coordinacion_id": None,
            })
    return decisiones, dict(Counter(r["estado"] for r in decisiones))
