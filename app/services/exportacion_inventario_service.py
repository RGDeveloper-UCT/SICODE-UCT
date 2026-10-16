"""Exportaciones seleccionables del Inventario Tecnológico."""
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, LongTable

# Claves públicas explícitas: no permitir atributos arbitrarios desde el formulario.
CAMPOS = {
    "numero_registro": ("No. registro", lambda e: e.numero_registro),
    "sicoin": ("SICOIN", lambda e: e.codigo),
    "tipo": ("Tipo de equipo", lambda e: e.tipo),
    "nombre": ("Nombre", lambda e: e.nombre),
    "marca": ("Marca", lambda e: e.marca),
    "modelo": ("Modelo", lambda e: e.modelo),
    "serie": ("No. serie", lambda e: e.serie),
    "ip_estatica": ("IP estática", lambda e: e.ip_estatica),
    "area": ("Coordinación", lambda e: e.area),
    "ubicacion": ("Ubicación", lambda e: e.ubicacion),
    "custodio": ("Custodio institucional", lambda e: e.nombre_responsable),
    "usuario_habitual": ("Usuario habitual", lambda e: e.nombre_usuario_habitual),
    "tipo_acceso": ("Tipo de acceso", lambda e: e.tipo_acceso),
    "acceso_institucional": ("Acceso institucional", lambda e: "Sí" if e.acceso_institucional else "No"),
    "estado": ("Estado", lambda e: e.estado),
    "boleta": ("Boleta de soporte", lambda e: e.numero_boleta),
    "observaciones": ("Observaciones", lambda e: e.observaciones),
    "grupo": ("Registro principal", lambda e: e.equipo_principal.numero_registro if e.equipo_principal else e.numero_registro),
}
FECHA_TZ = "America/Guatemala"


def seleccion(campos):
    if "todos" in campos:
        return list(CAMPOS)
    resultado = list(dict.fromkeys(campos))
    if not resultado or any(c not in CAMPOS for c in resultado):
        raise ValueError("Seleccione uno o más campos válidos.")
    return resultado


def datos(equipos, campos):
    return [[CAMPOS[c][1](e) for c in campos] for e in equipos]


def excel(equipos, campos, usuario, fecha):
    memoria = BytesIO()
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Inventario tecnológico"
    hoja.append(["SICODE-UCT | Inventario Tecnológico"])
    hoja.append([f"Origen: Coordinación / Panel de Carrete / Inventario Tecnológico"])
    hoja.append([f"Exportado: {fecha} | Usuario: {usuario}"])
    hoja.append([])
    hoja.append([CAMPOS[c][0] for c in campos])
    for fila in datos(equipos, campos):
        # Evitar inyección de fórmulas en datos de usuarios.
        hoja.append(["'" + str(v) if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) else (v if v is not None else "") for v in fila])
    hoja.freeze_panes = "A6"
    hoja.auto_filter.ref = f"A5:{get_column_letter(len(campos))}{max(5,hoja.max_row)}"
    for celda in hoja[5]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill("solid", fgColor="1F2937")
        celda.alignment = Alignment(wrap_text=True)
    for i, campo in enumerate(campos,1):
        hoja.column_dimensions[get_column_letter(i)].width = min(max(len(CAMPOS[campo][0])+3, 16), 38)
    libro.save(memoria)
    memoria.seek(0)
    return memoria


def pdf(equipos, campos, usuario, fecha):
    memoria = BytesIO()
    doc = SimpleDocTemplate(memoria, pagesize=landscape(A4), leftMargin=23,
                            rightMargin=23, topMargin=25, bottomMargin=38)
    ancho = landscape(A4)[0]-46
    estilo = ParagraphStyle("c", fontName="Helvetica", fontSize=6.5, leading=8.5)
    titulo = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=11, alignment=1, leading=14)
    pie = ParagraphStyle("p", fontName="Helvetica", fontSize=6.5, leading=9)
    def texto(v):
        return Paragraph(escape(str(v if v is not None else "")), estilo)
    def pie_pagina(canvas, documento):
        canvas.setFont("Helvetica", 7)
        canvas.drawString(23, 15, f"SICODE-UCT | Inventario Tecnológico | {fecha} | Exportado por: {usuario}")
        canvas.drawRightString(landscape(A4)[0]-23, 15, f"Página {documento.page}")
    cab = [Paragraph("<b>"+escape(CAMPOS[c][0])+"</b>", estilo) for c in campos]
    filas = [cab]+[[texto(v) for v in fila] for fila in datos(equipos, campos)]
    tabla = LongTable(filas, colWidths=[ancho/len(campos)]*len(campos), repeatRows=1, hAlign="LEFT")
    tabla.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#94A3B8")),
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#CBD5E1")),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),4),
        ("RIGHTPADDING",(0,0),(-1,-1),4),
        ("BOTTOMPADDING",(0,0),(-1,-1),4),
    ]))
    historia = [
        Paragraph("SICODE-UCT · MINISTERIO DE GOBERNACIÓN",titulo),
        Paragraph("INVENTARIO TECNOLÓGICO — REPORTE DE EXPORTACIÓN",titulo),
        Spacer(1,7),
        Paragraph(f"Origen: Coordinación / Panel de Carrete / Inventario Tecnológico. Fecha y hora: {escape(fecha)}. Usuario exportador: {escape(usuario)}.",pie),
        Paragraph(f"Registros: {len(equipos)}. Campos seleccionados: {len(campos)}.",pie),
        Spacer(1,10),tabla
    ]
    doc.build(historia,onFirstPage=pie_pagina,onLaterPages=pie_pagina)
    memoria.seek(0)
    return memoria
