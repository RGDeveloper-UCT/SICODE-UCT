"""Exportación configurable y auditada del directorio PERSONAL UCT."""
from io import BytesIO
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle

CAMPOS_PERSONAL = {
    "nombre": ("Nombre", lambda p: (p.nombre or "").upper()),
    "apellido": ("Apellido", lambda p: (p.apellido or "").upper()),
    "coordinacion": ("Coordinación", lambda p: p.coordinacion),
    "estado": ("Estado", lambda p: "Activo" if p.activo else "Inactivo"),
}
ORIGEN = "SICODE-UCT / Coordinación / Inventario Tecnológico / PERSONAL UCT"


def seleccionar_campos(valores):
    if "todos" in valores:
        return list(CAMPOS_PERSONAL)
    seleccion = list(dict.fromkeys(valores))
    if not seleccion or any(c not in CAMPOS_PERSONAL for c in seleccion):
        raise ValueError("Debe elegir uno o más campos válidos.")
    return seleccion


def _filas(personas, campos):
    return [[CAMPOS_PERSONAL[c][1](persona) for c in campos] for persona in personas]


def generar_excel(personas, campos, usuario, fecha):
    memoria = BytesIO()
    libro = Workbook()
    hoja = libro.active
    hoja.title = "PERSONAL UCT"
    hoja.append(["SICODE-UCT | PERSONAL UCT"])
    hoja.append(["Origen: " + ORIGEN])
    hoja.append([f"Fecha y hora: {fecha} | Exportado por: {usuario}"])
    hoja.append([])
    hoja.append([CAMPOS_PERSONAL[c][0] for c in campos])
    for fila in _filas(personas, campos):
        hoja.append([
            "'" + v if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@"))
            else (v if v is not None else "")
            for v in fila
        ])
    hoja.freeze_panes = "A6"
    hoja.auto_filter.ref = f"A5:{get_column_letter(len(campos))}{max(5, hoja.max_row)}"
    for celda in hoja[5]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F2937")
        celda.alignment = Alignment(wrap_text=True)
    for i, clave in enumerate(campos, 1):
        hoja.column_dimensions[get_column_letter(i)].width = 42 if clave == "coordinacion" else 24
    libro.save(memoria)
    memoria.seek(0)
    return memoria


def generar_pdf(personas, campos, usuario, fecha):
    memoria = BytesIO()
    ancho, _ = landscape(A4)
    doc = SimpleDocTemplate(memoria, pagesize=landscape(A4), leftMargin=27,
                            rightMargin=27, topMargin=26, bottomMargin=38)
    estilo = ParagraphStyle("celda", fontName="Helvetica", fontSize=8, leading=11)
    titulo = ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=12, leading=15, alignment=1)
    nota = ParagraphStyle("nota", fontName="Helvetica", fontSize=8, leading=11)
    def celda(valor):
        return Paragraph(escape(str(valor if valor is not None else "")), estilo)
    filas = [[celda(CAMPOS_PERSONAL[c][0]) for c in campos]]
    filas += [[celda(v) for v in fila] for fila in _filas(personas, campos)]
    tabla = LongTable(filas, colWidths=[(ancho - 54) / len(campos)] * len(campos),
                      repeatRows=1, hAlign="LEFT")
    tabla.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#94A3B8")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    def pie(canvas, document):
        canvas.setFont("Helvetica", 7)
        canvas.drawString(27, 17, f"SICODE-UCT | PERSONAL UCT | {fecha} | Usuario: {usuario}"[:115])
        canvas.drawRightString(ancho - 27, 17, f"Página {document.page}")
    historia = [
        Paragraph("MINISTERIO DE GOBERNACIÓN · UNIDAD DE CONTROL TELEMÁTICO", titulo),
        Paragraph("DIRECTORIO PERSONAL UCT — REPORTE DE EXPORTACIÓN", titulo),
        Spacer(1, 10),
        Paragraph(f"Origen: {escape(ORIGEN)}", nota),
        Paragraph(f"Fecha y hora: {escape(fecha)} | Exportado por: {escape(usuario)} | Registros: {len(personas)}", nota),
        Spacer(1, 10), tabla,
    ]
    doc.build(historia, onFirstPage=pie, onLaterPages=pie)
    memoria.seek(0)
    return memoria
