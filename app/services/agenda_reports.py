"""Exportadores autorizados para seguimiento de la Agenda UCT."""
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet


COLUMNAS = ("ID", "Tipo", "Actividad", "Estado", "Prioridad",
            "Responsable", "Inicio", "Vencimiento", "Referencia SP")


def celda_segura(valor):
    """Evita que Excel ejecute fórmulas provenientes de datos capturados."""
    texto = str(valor if valor is not None else "")
    if texto and (texto[0] in "=+-@" or texto[:1] in "\t\r"):
        return "'" + texto
    return texto


def filas_reporte(registros):
    for actividad in registros:
        yield (
            actividad.id, actividad.tipo, actividad.titulo,
            actividad.estado, actividad.prioridad,
            actividad.responsable.nombre if actividad.responsable else "",
            actividad.inicio.strftime("%d/%m/%Y %H:%M") if actividad.inicio else "",
            actividad.vencimiento.strftime("%d/%m/%Y %H:%M") if actividad.vencimiento else "",
            actividad.referencia_sp or "",
        )


def generar_xlsx(registros):
    wb = Workbook()
    ws = wb.active
    ws.title = "Agenda UCT"
    ws.append(COLUMNAS)
    for fila in filas_reporte(registros):
        ws.append([celda_segura(v) for v in fila])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for celda in ws[1]:
        celda.font = Font(color="FFFFFF", bold=True)
        celda.fill = PatternFill("solid", fgColor="102A4A")
        celda.alignment = Alignment(vertical="center")
    for letra, ancho in {
        "A": 9, "B": 17, "C": 48, "D": 20, "E": 14,
        "F": 30, "G": 20, "H": 20, "I": 18,
    }.items():
        ws.column_dimensions[letra].width = ancho
    ws.sheet_view.showGridLines = False
    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida


def generar_pdf(registros):
    salida = BytesIO()
    documento = SimpleDocTemplate(salida, pagesize=landscape(A4), rightMargin=26,
                                  leftMargin=26, topMargin=28, bottomMargin=27)
    estilos = getSampleStyleSheet()
    normal = estilos["BodyText"]
    normal.fontSize = 7
    normal.leading = 9
    titulo = Paragraph("SICODE-UCT · Reporte de actividades", estilos["Heading2"])
    contenido = [titulo, Spacer(1, 13)]
    tabla = [[Paragraph(escape(str(c)), normal) for c in COLUMNAS[:8]]]
    for fila in filas_reporte(registros):
        tabla.append([Paragraph(escape(str(col)[:85]),normal) for col in fila[:8]])
    celdas = Table(tabla, colWidths=[33,67,215,73,52,118,85,85], repeatRows=1, hAlign="LEFT")
    celdas.setStyle(TableStyle([
        ("BACKGROUND", (0,0),(-1,0),colors.HexColor("#173457")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#d4dce8")),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f1f5fa")]),
        ("TOPPADDING",(0,0),(-1,-1),5),
        ("BOTTOMPADDING",(0,0),(-1,-1),5),
    ]))
    contenido.append(celdas)
    documento.build(contenido)
    salida.seek(0)
    return salida
